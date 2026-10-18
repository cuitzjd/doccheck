"""Loader and repair must agree on DOCX paragraph indexes."""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx

from doccheck.docloader import load
from doccheck.repair import _iter_paragraphs, _paragraph_by_index


with tempfile.TemporaryDirectory(prefix="doccheck_paragraphs_") as td:
    path = os.path.join(td, "paragraphs.docx")
    document = docx.Document()
    document.add_paragraph("正文第一段")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].text = "表格段落"
    document.add_paragraph("正文第二段")
    document.save(path)

    loaded = load(path)
    raw = docx.Document(path)
    repaired = list(_iter_paragraphs(raw))
    expected = ["正文第一段", "表格段落", "正文第二段"]

    assert [p.text for p in loaded.paragraphs] == expected
    assert [p.text for p in repaired] == expected
    assert [_paragraph_by_index(raw, i).text for i in range(len(expected))] == expected
print("DOCX paragraph enumeration tests passed")
