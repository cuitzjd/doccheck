"""指引册生成（GUIDE-01/02）：从已交付文档自动抽取章节标题与摘要生成 docx 初稿。

抽不到的内容留【待补充】占位，由用户补全（规格书 3.7：生成初稿）。
"""
import re

from .models import UnifiedDoc

_PLACEHOLDER = "【待补充】"


def _headings(doc: UnifiedDoc, max_level_chars=40):
    return [p.text for p in doc.paragraphs if p.is_heading and len(p.text) < max_level_chars]


def _first_body_after_heading(doc: UnifiedDoc, maxlen=120) -> str:
    """首个标题之后的第一段正文作为摘要。"""
    seen_heading = False
    for p in doc.paragraphs:
        if p.is_heading:
            seen_heading = True
            continue
        if seen_heading and not p.is_table_cell and len(p.text) >= 20:
            return p.text[:maxlen] + ("…" if len(p.text) > maxlen else "")
    return ""


def _section_summary(doc: UnifiedDoc, kw: str, maxlen=200) -> str:
    """抽取包含 kw 的标题章节的首段文字。"""
    hit = False
    for p in doc.paragraphs:
        if p.is_heading:
            hit = kw in p.text
            continue
        if hit and not p.is_table_cell and len(p.text) >= 20:
            return p.text[:maxlen] + ("…" if len(p.text) > maxlen else "")
    return ""


def generate_doc_guide(docs: list, out_path: str) -> str:
    """GUIDE-02 文档指引册：逐份介绍文档名称/类型/页数/章节结构/主要内容。"""
    import docx
    document = docx.Document()
    document.add_heading("文档指引册", 0)
    document.add_paragraph("本册对本次交付各文档的内容作简要说明，帮助读者快速找到目标文档。")
    for d in docs:
        document.add_heading(d.name, level=1)
        tbl = document.add_table(rows=5, cols=2)
        tbl.style = "Table Grid"
        heads = _headings(d)
        rows = [("文档名称", d.name),
                ("文档类型", d.doc_type),
                ("页数", str(d.page_count) if d.page_count else "未知"),
                ("章节结构", "\n".join(heads[:40]) or _PLACEHOLDER),
                ("主要内容", _first_body_after_heading(d) or _PLACEHOLDER)]
        for row, (k, v) in zip(tbl.rows, rows):
            row.cells[0].text = k
            row.cells[1].text = v
    document.save(out_path)
    return out_path


def generate_brief_guide(docs: list, contract, out_path: str) -> str:
    """GUIDE-01 项目汇报材料指引册骨架。

    固定章节：验收范围/项目背景/建设对象/建设成效/资金情况/软件开发业务介绍
    + 专家角色导读索引 + 册末附件功能点清单。可抽取内容自动填入，其余【待补充】。
    """
    import docx
    document = docx.Document()
    proj = getattr(contract, "project_name", "") or _PLACEHOLDER
    document.add_heading("项目汇报材料指引册", 0)
    document.add_paragraph(f"项目名称：{proj}")

    def find_doc(*kws):
        for d in docs:
            if any(k in d.name or k in d.doc_type for k in kws):
                return d
        return None

    # 一、验收范围
    document.add_heading("一、验收范围", level=1)
    v = ""
    d = find_doc("阶段性服务总结报告")
    if d:
        v = _section_summary(d, "服务内容")
    document.add_paragraph(v or _PLACEHOLDER)

    # 二、项目背景
    document.add_heading("二、项目背景", level=1)
    v = ""
    d = find_doc("需求规格说明书", "项目实施方案")
    if d:
        v = _section_summary(d, "项目背景") or _section_summary(d, "背景")
    document.add_paragraph(v or _PLACEHOLDER)

    # 三、建设对象
    document.add_heading("三、建设对象", level=1)
    subs = []
    d = find_doc("需求规格说明书", "技术设计方案")
    if d:
        subs = [h for h in _headings(d) if "子系统" in h][:12]
    if subs:
        for s in subs:
            document.add_paragraph(s, style="List Bullet")
    else:
        document.add_paragraph(_PLACEHOLDER)

    # 四、建设成效
    document.add_heading("四、建设成效", level=1)
    v = ""
    d = find_doc("阶段性服务总结报告")
    if d:
        v = _section_summary(d, "成效") or _section_summary(d, "成果")
    document.add_paragraph(v or _PLACEHOLDER)

    # 五、资金情况
    document.add_heading("五、资金情况", level=1)
    amount = ""
    if getattr(contract, "amount_cn", "") or getattr(contract, "amount_number", ""):
        amount = f"合同金额：{contract.amount_cn or ''}（{contract.amount_number} 元）"
        document.add_paragraph(amount)
        for i, t in enumerate(getattr(contract, "payment_terms", []) or [], 1):
            document.add_paragraph(f"付款节点{i}：{t}", style="List Bullet")
    else:
        document.add_paragraph(_PLACEHOLDER)

    # 六、软件开发业务介绍
    document.add_heading("六、软件开发业务介绍", level=1)
    v = ""
    d = find_doc("系统内部测试报告")
    if d:
        v = _section_summary(d, "测试")[:200]
    document.add_paragraph(v or _PLACEHOLDER)

    # 专家导读索引
    document.add_heading("附：专家关注点导读索引", level=1)
    tbl = document.add_table(rows=4, cols=2)
    tbl.style = "Table Grid"
    for row, (role, where) in zip(tbl.rows, [
            ("财务专家", "五、资金情况（合同金额、付款条件）"),
            ("业务专家", "二、项目背景；三、建设对象；四、建设成效"),
            ("技术/软开专家", "六、软件开发业务介绍；附件：功能点清单"),
            ("其他", _PLACEHOLDER)]):
        row.cells[0].text = role
        row.cells[1].text = where

    # 册末附件：功能点清单（占位，具体功能点不入正文）
    document.add_page_break()
    document.add_heading("附件：功能点完成情况清单", level=1)
    document.add_paragraph("（具体功能点统一置于本附件，不放入正文）")
    tbl = document.add_table(rows=2, cols=3)
    tbl.style = "Table Grid"
    for c, t in zip(tbl.rows[0].cells, ("序号", "功能点", "完成情况")):
        c.text = t
    tbl.rows[1].cells[1].text = _PLACEHOLDER

    document.save(out_path)
    return out_path
