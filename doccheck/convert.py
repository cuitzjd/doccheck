"""PDF↔Word 转换（CONVERT-01/02/04/05）。

转换件只用于内容类检查与修复，不作为版式判定依据（规格书 3.9 说明）。

后端优先级：
1. WPS wpscli（/Applications/wpsoffice.app，保真度最高）。注意 WPS 为沙盒应用，
   只能读写用户授权目录（实测 ~/Downloads 可用，/tmp 与项目目录被拒绝），
   因此一律经 ~/Downloads/doccheck/ 中转；
2. 内置实现（PyMuPDF 文本层重建 / textutil / Vision OCR）；
3. LibreOffice（word2pdf 兜底）。
"""
import os
import re
import shutil
import subprocess

from . import docloader, ocr

_WPSCLI = "/Applications/wpsoffice.app/Contents/MacOS/wpscli"
_WPS_STAGE = os.path.expanduser("~/Downloads/doccheck")


def _wpscli() -> str:
    return _WPSCLI if os.path.exists(_WPSCLI) else ""


def _wps_run(subcmd: str, src: str, out_path: str) -> tuple:
    """经中转目录调用 wpscli 转换。返回 (ok, msg)。"""
    os.makedirs(_WPS_STAGE, exist_ok=True)
    stage_src = os.path.join(_WPS_STAGE, os.path.basename(src))
    stage_out = os.path.join(_WPS_STAGE, os.path.basename(out_path))
    if os.path.abspath(src) != os.path.abspath(stage_src):
        shutil.copy2(src, stage_src)
    for p in (stage_out,):
        if os.path.exists(p):
            os.remove(p)
    r = subprocess.run([_WPSCLI, subcmd, stage_src, "-o", stage_out, "--json"],
                       capture_output=True, text=True, timeout=3600)
    if not os.path.exists(stage_out):
        err = (r.stdout + r.stderr).strip().splitlines()
        err = [l for l in err if "QMutex" not in l]
        return False, (err[-1][:200] if err else f"wpscli {subcmd} 失败")
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    if os.path.abspath(stage_out) != os.path.abspath(out_path):
        shutil.copy2(stage_out, out_path)
    for p in (stage_out, stage_src):
        try:
            if os.path.exists(p) and os.path.abspath(p) != os.path.abspath(src):
                os.remove(p)
        except OSError:
            pass
    return True, ""


def pdf_to_docx(pdf_path: str, out_path: str) -> tuple:
    """文字版 PDF → docx。优先 WPS 本地引擎（保留表格/版式，保真度高）；
    WPS 不可用时退回内置实现（按字号分档重建标题层级，表格按行重建）。"""
    import fitz
    pdf = fitz.open(pdf_path)
    total_chars = sum(len(pdf[i].get_text()) for i in range(min(5, pdf.page_count)))
    n_pages = pdf.page_count
    pdf.close()
    if n_pages and total_chars < 100:
        return False, "PDF 无文本层（扫描件），请改用 pdf_ocr_to_docx", ""
    if _wpscli():
        ok, err = _wps_run("pdf2word", pdf_path, out_path)
        if ok:
            return True, f"转换完成（WPS 本地引擎）：{n_pages} 页 → {out_path}", out_path
        # 失败则继续走内置实现
    return _pdf_to_docx_builtin(pdf_path, out_path, n_pages)


def _pdf_to_docx_builtin(pdf_path: str, out_path: str, n_pages: int = 0) -> tuple:
    import fitz
    import docx
    pdf = fitz.open(pdf_path)
    n_pages = n_pages or pdf.page_count
    document = docx.Document()
    n_chars = 0
    no_text_pages = []
    for pno in range(pdf.page_count):
        pg = pdf[pno]
        if not pg.get_text().strip():
            no_text_pages.append(pno + 1)
            continue
        d = pg.get_text("dict")
        sizes = sorted({round(s["size"], 1) for b in d["blocks"]
                        for l in b.get("lines", []) for s in l["spans"]
                        if s["text"].strip()}, reverse=True)
        # 字号分档：最大两档视为标题（正文通常为第 3 档及以下）
        h1 = sizes[0] if len(sizes) >= 2 else None
        h2 = sizes[1] if len(sizes) >= 3 else None
        body_size = sizes[2] if len(sizes) >= 3 else (sizes[-1] if sizes else 0)
        for b in d["blocks"]:
            for l in b.get("lines", []):
                text = "".join(s["text"] for s in l["spans"]).strip()
                if not text:
                    continue
                n_chars += len(text)
                size = max((s["size"] for s in l["spans"]), default=0)
                bold = any("Bold" in s["font"] or "bold" in s["font"]
                           for s in l["spans"])
                if h1 and size >= h1 and size > body_size and len(text) < 60:
                    document.add_heading(text, level=1)
                elif h2 and size >= h2 and size > body_size and len(text) < 60:
                    document.add_heading(text, level=2)
                elif bold and size > body_size and len(text) < 60:
                    document.add_heading(text, level=3)
                else:
                    document.add_paragraph(text)
    n_pages = pdf.page_count
    pdf.close()
    document.save(out_path)
    msg = f"转换完成：{n_pages} 页 → {out_path}（约 {n_chars} 字；表格按行重建，版式以原件为准）"
    if no_text_pages:
        msg += f"；注意：第 {'、'.join(map(str, no_text_pages))} 页无文本层（图片页），如需内容请用 pdf_ocr_to_docx"
    return True, msg, out_path


def pdf_ocr_to_docx(pdf_path: str, out_path: str) -> tuple:
    """扫描件 OCR → docx（本地 OCR，见 ocr.py）。

    注意：WPS 的扫描件转换（--scanned）与 photo2word 均为云端 OCR，
    涉密文档不可用，此处只走本地后端（tesseract / macOS Vision）。"""
    import docx
    if not ocr.available():
        return False, "本机无可用 OCR 后端（tesseract/Vision 均不可用），请手工录入", ""
    text = ocr.ocr_pdf(pdf_path)
    if not text.strip():
        return False, "OCR 未识别出文本", ""
    document = docx.Document()
    n = 0
    for line in text.splitlines():
        line = line.strip()
        if not line or re.match(r"^===== 第\d+页 =====$", line):
            continue
        document.add_paragraph(line)
        n += len(line)
    document.save(out_path)
    return True, f"OCR 转换完成（后端 {ocr.available()}，约 {n} 字）：关键字段须人工核对后方可入库", out_path


def doc_to_docx(path: str, out_path: str = "") -> tuple:
    """老式 .doc → .docx（macOS textutil，经 docloader）。"""
    try:
        tmp = docloader.doc_to_docx(path)
    except Exception as e:
        return False, f"textutil 转换失败: {e}", ""
    if out_path and os.path.abspath(tmp) != os.path.abspath(out_path):
        shutil.copy2(tmp, out_path)
        tmp = out_path
    return True, f"转换完成：{tmp}（页数/段落数请抽查校验）", tmp


def _soffice() -> str:
    for c in (shutil.which("soffice"), shutil.which("libreoffice"),
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if c and os.path.exists(c):
            return c
    return ""


def word_to_pdf(path: str, out_path: str) -> tuple:
    """Word → PDF。优先 WPS wpscli word2pdf（本地，支持 .doc/.docx）；
    其次 LibreOffice；都不可用则明确返回失败，不伪造成功。"""
    if _wpscli():
        ok, err = _wps_run("word2pdf", path, out_path)
        if ok:
            return True, f"导出完成（WPS 本地引擎）：{out_path}（请校验页数与源文档一致）", out_path
        # 失败继续尝试 LibreOffice
    exe = _soffice()
    if not exe:
        return False, "需要安装 WPS Office 或 LibreOffice 后方可导出 PDF", ""
    outdir = os.path.dirname(os.path.abspath(out_path)) or "."
    r = subprocess.run([exe, "--headless", "--convert-to", "pdf",
                        "--outdir", outdir, path],
                       capture_output=True, text=True, timeout=600)
    produced = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    if r.returncode != 0 or not os.path.exists(produced):
        return False, f"LibreOffice 转换失败: {r.stderr.strip()[:200]}", ""
    if os.path.abspath(produced) != os.path.abspath(out_path):
        shutil.move(produced, out_path)
    return True, f"导出完成：{out_path}（请校验页数与源文档一致）", out_path
