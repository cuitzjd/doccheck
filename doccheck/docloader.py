"""文档加载器：把 docx / doc / pdf 统一解析为 UnifiedDoc。

要点（来自样例实测）：
- docx 正文字体常继承自样式（run 级为 None），必须沿 run → 段落样式 → 文档默认 逐级解析；
- docx 页码是页脚 PAGE 域代码，页码类检查依赖配对的盖章 PDF；
- pdf 封面为图片页（无文本），文档类型靠文件名前缀识别；
- .doc 老格式先经 textutil 转 .docx（macOS 自带）。
"""
import os
import re
import subprocess
import tempfile

from .models import Paragraph, UnifiedDoc
from .docx_paragraphs import iter_docx_paragraphs

DOC_TYPE_PREFIX = re.compile(r"^【(?:\d+-)?([^】]+)】")


def detect_doc_type(filename: str) -> str:
    """从文件名识别文档类型：【04-需求规格说明书】xxx.pdf → 需求规格说明书"""
    m = DOC_TYPE_PREFIX.match(os.path.basename(filename))
    if m:
        return m.group(1).strip()
    name = os.path.basename(filename)
    for kw in ["指引册", "阶段性服务总结报告", "月进度计划报告", "合同"]:
        if kw in name:
            return "中期验收指引册" if kw == "指引册" else kw
    return "未识别"


def load(path: str, paired_pdf: str = "") -> UnifiedDoc:
    ext = os.path.splitext(path)[1].lower()
    orig_path, converted = path, False
    if ext == ".doc":
        path = doc_to_docx(path)  # 转换后按 docx 处理
        ext, converted = ".docx", True
    doc = UnifiedDoc(path=path, name=os.path.basename(orig_path),  # 名称始终用原文件名
                     doc_type=detect_doc_type(orig_path), fmt=ext.lstrip("."),
                     paired_pdf=paired_pdf)
    if converted:
        doc.meta["original_path"] = orig_path
        doc.meta["converted_from_doc"] = True
    if ext == ".docx":
        _load_docx(doc)
    elif ext == ".pdf":
        _load_pdf(doc)
    else:
        raise ValueError(f"不支持的格式: {path}")
    return doc


# ---------------- docx ----------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _resolve_fonts(d, para, run):
    """沿 run → 段落样式 → 文档默认 解析生效字体/字号。"""
    rpr = run.font.element.rPr
    font_cn = font_en = None
    size = run.font.size.pt if run.font.size else None
    if rpr is not None and rpr.rFonts is not None:
        font_cn = rpr.rFonts.get(_W + "eastAsia")
        font_en = rpr.rFonts.get(_W + "ascii")
    st = para.style
    while st is not None:
        if font_cn is None or font_en is None:
            try:
                rf = st.element.rPr.rFonts if st.element.rPr is not None else None
            except AttributeError:
                rf = None
            if rf is not None:
                font_cn = font_cn or rf.get(_W + "eastAsia")
                font_en = font_en or rf.get(_W + "ascii")
        if size is None and getattr(st.font, "size", None):
            size = st.font.size.pt
        st = getattr(st, "base_style", None)
    # 文档默认
    defaults = d.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")
    if defaults:
        ddf = defaults[0]
        rf = ddf.rFonts
        if rf is not None:
            font_cn = font_cn or rf.get(_W + "eastAsia")
            font_en = font_en or rf.get(_W + "ascii")
        if size is None and ddf.sz is not None:
            raw_size = ddf.sz.get(_W + "val")
            try:
                size = int(raw_size) / 2
            except (TypeError, ValueError):
                pass
    return font_cn or "", font_en or "", size or 0.0


def _xml_bool(parent, name: str, default: bool = False) -> bool:
    """读取 w:pPr 下的布尔属性，兼容 w:val=0/false 和无 val。"""
    if parent is None:
        return default
    element = parent.find(_W + name)
    if element is None:
        return default
    value = element.get(_W + "val")
    return value is None or str(value).lower() not in {"0", "false", "off", "no"}


def _paragraph_properties(para):
    """提取编号、标题大纲和分页等不会体现在段落文本里的 Word 结构。"""
    ppr = para._p.pPr
    num_id, num_level = "", -1
    if ppr is not None:
        num_pr = ppr.find(_W + "numPr")
        if num_pr is not None:
            num = num_pr.find(_W + "numId")
            level = num_pr.find(_W + "ilvl")
            if num is not None:
                num_id = num.get(_W + "val", "")
            if level is not None:
                try:
                    num_level = int(level.get(_W + "val", "-1"))
                except (TypeError, ValueError):
                    num_level = -1
    outline = -1
    if ppr is not None:
        outline_el = ppr.find(_W + "outlineLvl")
        if outline_el is not None:
            try:
                outline = int(outline_el.get(_W + "val", "-1"))
            except (TypeError, ValueError):
                outline = -1
    return {
        "num_id": num_id,
        "num_level": num_level,
        "is_list_item": bool(num_id and num_id != "0"),
        "keep_with_next": _xml_bool(ppr, "keepNext"),
        "page_break_before": _xml_bool(ppr, "pageBreakBefore"),
        "outline_level": outline,
    }


def _effective_bold(para):
    """取 run 直设或样式继承后的粗体，避免标题样式被当成普通正文。"""
    for run in para.runs:
        if run.text.strip() and run.font.bold is True:
            return True
    style = para.style
    while style is not None:
        if getattr(style.font, "bold", None) is True:
            return True
        style = getattr(style, "base_style", None)
    return False


def _load_docx(doc: UnifiedDoc):
    import docx
    d = docx.Document(doc.path)
    page = 1
    idx = 0
    for item in iter_docx_paragraphs(d):
        idx = _emit_docx_para(doc, d, item.paragraph, idx, page, item.in_table)
        page += item.page_breaks
    doc.page_count = page  # 估算值；精确页数以配对 PDF 为准
    # 分节/页脚域代码信息（供 PAGE-03/05 结构检查）
    doc.meta["sections"] = len(d.sections)
    doc.meta["sect_page_restarts"] = [
        (el.get(_W + "start") if el is not None else None)
        for s in d.sections
        for el in [s._sectPr.find(_W + "pgNumType")]
    ]
    doc.meta["footer_has_page_field"] = any(
        s.footer._element.xpath(".//w:instrText[contains(., 'PAGE')]") or
        s.footer._element.xpath(".//w:fldSimple[contains(@w:instr, 'PAGE')]")
        for s in d.sections
    )
    doc.meta["footer_page_field_by_section"] = [
        bool(s.footer._element.xpath(".//w:instrText[contains(., 'PAGE')]") or
             s.footer._element.xpath(".//w:fldSimple[contains(@w:instr, 'PAGE')]"))
        for s in d.sections
    ]
    doc.meta["footer_texts"] = [
        "".join(p.text for p in s.footer.paragraphs).strip() for s in d.sections
    ]


def _emit_docx_para(doc, d, para, idx, page, in_table=False):
    text = para.text.strip()
    if not text:
        return idx
    style = para.style.name if para.style is not None else ""
    properties = _paragraph_properties(para)
    is_heading = (style.lower().startswith(("heading", "title", "subtitle")) or
                  style.startswith(("标题", "副标题")) or
                  properties["outline_level"] >= 0)
    font_cn = font_en = ""
    size = 0.0
    bold = _effective_bold(para)
    for run in para.runs:
        if run.text.strip():
            font_cn, font_en, size = _resolve_fonts(d, para, run)
            break
    ppr = para.paragraph_format
    doc.paragraphs.append(Paragraph(
        index=idx, text=text, page=page, style=style,
        font_cn=font_cn, font_en=font_en, font_size=size, bold=bold,
        is_heading=is_heading, is_table_cell=in_table,
        alignment=str(ppr.alignment or ""),
        first_line_indent=ppr.first_line_indent.pt if ppr.first_line_indent else None,
        line_spacing=ppr.line_spacing if isinstance(ppr.line_spacing, float) else None,
        num_id=properties["num_id"], num_level=properties["num_level"],
        is_list_item=properties["is_list_item"],
        keep_with_next=properties["keep_with_next"],
        page_break_before=properties["page_break_before"],
        outline_level=properties["outline_level"],
    ))
    return idx + 1


# ---------------- pdf ----------------

# 页脚页码：标准"第 X 页"；部分盖章 PDF 页脚文本层缺"页"字（仅剩"第 10 "），
# 且存在"第 X 页 共 N 页"格式，故分严格/宽松两级匹配
_FOOTER_PAGENO = re.compile(r"第\s*(\d+)\s*页")
_FOOTER_PAGENO_LOOSE = re.compile(r"^第\s*(\d+)\s*$", re.M)
_FOOTER_PAGENO_BARE = re.compile(r"^\s*(\d{1,4})\s*$", re.M)  # 裸数字页脚（如代码编写规范）
_FOOTER_FRACTION = re.compile(
    r"(?<!\d)(\d{1,4})\s*[/／]\s*(\d{1,4})\s*页?(?!\d)"
)
_FOOTER_TOTAL = re.compile(r"共\s*(\d{1,4})\s*页")
_TOC_LINE = re.compile(r"^(.*?)[.·…︙]{2,}\s*(\d+)\s*$")


def _load_pdf(doc: UnifiedDoc):
    import fitz
    pdf = fitz.open(doc.path)
    doc.page_count = pdf.page_count
    idx = 0
    in_toc = False
    toc_pages = 0
    bare_last = None  # 裸数字页脚的连续追踪
    for pno in range(pdf.page_count):
        pg = pdf[pno]
        text = pg.get_text()
        if not text.strip():
            continue  # 图片页（封面/扫描页）
        lines = [l.strip() for l in text.splitlines() if l.strip()]
        # 页脚页码：取页面底部区域内的"第 X 页"（文本提取顺序不可靠）。
        # 实测部分文档页脚块顶边位于 91.9% 页高处，故按块纵向中心 >90% 判定
        page_h = pg.rect.height
        foot_blocks = [b[4] for b in pg.get_text("blocks")
                       if (b[1] + b[3]) / 2 > page_h * 0.90]
        footer_text = "\n".join(foot_blocks)
        fraction = None
        for txt in foot_blocks:
            fraction = _FOOTER_FRACTION.search(txt)
            if fraction:
                break
        if fraction:
            current, total = int(fraction.group(1)), int(fraction.group(2))
            doc.footer_page_numbers.append((pno + 1, current))
            doc.footer_page_totals.append((pno + 1, total))
            doc.footer_page_fractions.append((pno + 1, current, total))
            bare_last = current
        else:
            total_match = _FOOTER_TOTAL.search(footer_text)
            if total_match:
                doc.footer_page_totals.append((pno + 1, int(total_match.group(1))))
        m = None
        if not fraction:
            for pat in (_FOOTER_PAGENO, _FOOTER_PAGENO_LOOSE):
                for txt in foot_blocks:
                    m = pat.search(txt)
                    if m:
                        break
                if m:
                    break
        if m:
            doc.footer_page_numbers.append((pno + 1, int(m.group(1))))
            bare_last = int(m.group(1))
        elif not fraction:
            # 裸数字页脚兜底：部分 PDF 文本层缺字（实测出现"15"只剩"1"），
            # 仅接受与上一页严格连续的裸数字，避免把页底表格数字当页码
            n = None
            for txt in foot_blocks:
                bm = _FOOTER_PAGENO_BARE.search(txt)
                if bm:
                    n = int(bm.group(1))
                    break
            if n is not None and (bare_last is None and n == 1
                                  or bare_last is not None and n == bare_last + 1):
                doc.footer_page_numbers.append((pno + 1, n))
                bare_last = n
        # 目录识别：出现"目 录/目录"标题后，随后 1~3 页视为目录
        joined = "".join(lines[:3])
        if re.match(r"^目\s*录$", joined) or joined.startswith("目录"):
            in_toc, toc_pages = True, 3
        elif in_toc:
            toc_pages -= 1
            if toc_pages <= 0 or not any(_TOC_LINE.match(l) for l in lines):
                in_toc = False
        for line in lines:
            if in_toc:
                tm = _TOC_LINE.match(line)
                if tm:
                    doc.toc_entries.append((tm.group(1).strip(), int(tm.group(2))))
            doc.paragraphs.append(Paragraph(index=idx, text=line, page=pno + 1))
            idx += 1
    if doc.paragraphs:
        doc.cover_text = doc.paragraphs[0].text if doc.paragraphs[0].page == 1 else ""
    pdf.close()


# ---------------- doc 转换（macOS textutil） ----------------

def doc_to_docx(path: str) -> str:
    out = os.path.join(tempfile.mkdtemp(prefix="doccheck_"),
                       os.path.splitext(os.path.basename(path))[0] + ".docx")
    subprocess.run(["textutil", "-convert", "docx", path, "-output", out],
                   check=True, capture_output=True)
    return out
